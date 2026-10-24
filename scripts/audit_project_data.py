
from pathlib import Path
import csv
import importlib.util

PROJECT = Path(r"D:\projects\energy_forecasting_io")
FIGARO = PROJECT / "data" / "FIGARO database"
ELECTRICITY = PROJECT / "data" / "raw" / "germany_electricity_hourly.csv"

print("=" * 72)
print("ENERGY FORECASTING IO — DATA AUDIT")
print("=" * 72)

# 1. Inventory: list project data files without modifying them
print("\n[1] DATA FILE INVENTORY")

data_dir = PROJECT / "data"
files = sorted(
    p for p in data_dir.rglob("*")
    if p.is_file()
)

for p in files:
    try:
        size_mb = p.stat().st_size / (1024 * 1024)
        print(f"{size_mb:9.2f} MB | {p.relative_to(PROJECT)}")
    except OSError as exc:
        print("Cannot inspect:", p, exc)

# 2. Hourly electricity data audit
print("\n[2] HOURLY ELECTRICITY DATA")

if not ELECTRICITY.exists():
    print("ERROR: Electricity CSV not found:", ELECTRICITY)
else:
    try:
        import pandas as pd

        df = pd.read_csv(ELECTRICITY, low_memory=False)

        print("Rows:", len(df))
        print("Columns:", len(df.columns))
        print("Column names:", list(df.columns))
        print("\nMissing values by column:")
        print(df.isna().sum().to_string())

        # Identify likely timestamp columns
        time_cols = [
            c for c in df.columns
            if any(word in c.lower()
                   for word in ("time", "date", "timestamp"))
        ]

        for col in time_cols:
            parsed = pd.to_datetime(df[col], errors="coerce", utc=True)
            valid = parsed.dropna()
            print(f"\nTimestamp column: {col}")
            print("Unparseable:", int(parsed.isna().sum()))
            if not valid.empty:
                print("First:", valid.min())
                print("Last: ", valid.max())
                print("Unique timestamps:", int(valid.nunique()))
                print("Duplicate timestamps:",
                      int(parsed.duplicated().sum()))

        numeric_cols = df.select_dtypes(include="number").columns
        if len(numeric_cols):
            print("\nNumeric summary:")
            print(df[numeric_cols].describe().round(3).to_string())

    except ImportError:
        print("pandas is unavailable in the project environment.")
    except Exception as exc:
        print("Electricity audit error:", repr(exc))

# 3. FIGARO CSV inspection
print("\n[3] FIGARO MATRIX CSV FILES")

csv_files = sorted(FIGARO.rglob("*.csv")) if FIGARO.exists() else []

if not csv_files:
    print("No CSV files found under:", FIGARO)
else:
    print("CSV files found:", len(csv_files))

    for path in csv_files:
        print("\n" + "-" * 65)
        print("FILE:", path.relative_to(PROJECT))
        print("SIZE MB:", round(path.stat().st_size / (1024 * 1024), 2))

        try:
            with path.open(
                "r", encoding="utf-8-sig", errors="replace", newline=""
            ) as f:
                sample = [f.readline() for _ in range(5)]

            for i, line in enumerate(sample, start=1):
                if not line:
                    break
                # Limit output for very wide matrix rows
                line = line.rstrip("\r\n")
                print(f"Line {i}: {line[:350]}")

            # Estimate delimiter from the first non-empty line
            first = next((x for x in sample if x.strip()), "")
            try:
                dialect = csv.Sniffer().sniff(first, delimiters=",;\t|")
                print("Detected delimiter:", repr(dialect.delimiter))
            except csv.Error:
                print("Delimiter could not be detected automatically.")

        except Exception as exc:
            print("CSV inspection error:", repr(exc))

# 4. Excel workbook inventory and sheet names
print("\n[4] EXCEL WORKBOOKS")

excel_files = sorted(
    p for p in data_dir.rglob("*")
    if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
)

if not excel_files:
    print("No .xlsx or .xlsm files found.")
elif importlib.util.find_spec("openpyxl") is None:
    print("Excel files found, but openpyxl is not installed.")
    for p in excel_files:
        print(p.relative_to(PROJECT))
else:
    from openpyxl import load_workbook

    for path in excel_files:
        print("\nWorkbook:", path.relative_to(PROJECT))
        try:
            wb = load_workbook(
                path, read_only=True, data_only=True
            )
            for ws in wb.worksheets:
                print(
                    f"  Sheet: {ws.title!r} | "
                    f"Rows: {ws.max_row} | Columns: {ws.max_column}"
                )
            wb.close()
        except Exception as exc:
            print("Workbook inspection error:", repr(exc))

print("\n" + "=" * 72)
print("AUDIT FINISHED — no source files were modified.")
print("=" * 72)
